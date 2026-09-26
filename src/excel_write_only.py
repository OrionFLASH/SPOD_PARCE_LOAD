# -*- coding: utf-8 -*-
"""
Потоковая запись книги Excel (openpyxl write_only) по готовому плану листов — PERF-07, этап 2.

Включается флагом ``performance.excel_writer: "write_only"`` (по умолчанию ``"openpyxl"`` — прежний
путь pandas.to_excel + оформление). План (стили заголовка и колонок, ширины, закрепление, фильтр)
строит src/main_impl.py так, чтобы книга совпадала с прежней; здесь — только запись:
строки идут в файл потоком, объекты ячеек не накапливаются в памяти.

Значения переводятся в ячейки так же, как это делает pandas ExcelWriter (см. excel_value_with_format).
Только stdlib + openpyxl/pandas из Anaconda 3.12.
"""

from __future__ import annotations

import datetime as dt
from copy import copy
from dataclasses import dataclass, field
from typing import Any, Callable, Dict, List, Optional, Tuple

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.utils import get_column_letter

# Форматы дат pandas ExcelWriter по умолчанию
PANDAS_DATETIME_FORMAT = "YYYY-MM-DD HH:MM:SS"
PANDAS_DATE_FORMAT = "YYYY-MM-DD"


def excel_value_with_format(val: Any) -> Tuple[Any, Optional[str]]:
    """
    Значение ячейки и формат — как pandas ExcelWriter (ExcelFormatter._format_value +
    ExcelWriter._value_with_fmt): пропуск → "", целые → int, дробные → float (inf → "inf"),
    datetime/date — с форматом даты, timedelta — доля суток с форматом "0", остальное — str.
    """
    if val is None:
        return "", None
    if isinstance(val, str):
        return val, None
    try:
        if pd.isna(val):
            return "", None
    except (TypeError, ValueError):
        pass
    if isinstance(val, (bool, np.bool_)):
        return bool(val), None
    if isinstance(val, (int, np.integer)):
        return int(val), None
    if isinstance(val, (float, np.floating)):
        f = float(val)
        if f == float("inf"):
            return "inf", None
        if f == float("-inf"):
            return "-inf", None
        return f, None
    if isinstance(val, dt.datetime):
        return val, PANDAS_DATETIME_FORMAT
    if isinstance(val, dt.date):
        return val, PANDAS_DATE_FORMAT
    if isinstance(val, (dt.timedelta, pd.Timedelta)):
        return pd.Timedelta(val).total_seconds() / 86400, "0"
    return str(val), None


@dataclass
class CellStyle:
    """Оформление ячейки; None — атрибут не задаётся (остаётся по умолчанию)."""

    font: Optional[Font] = None
    fill: Optional[PatternFill] = None
    border: Optional[Border] = None
    alignment: Optional[Alignment] = None
    number_format: Optional[str] = None

    def is_empty(self) -> bool:
        return all(v is None for v in (self.font, self.fill, self.border, self.alignment, self.number_format))


@dataclass
class ColumnPlan:
    """Данные колонки: оформление и преобразование значения."""

    style: CellStyle = field(default_factory=CellStyle)
    # Формат числа/даты из COLUMN_FORMATS (перекрывает формат даты pandas); None — формат pandas
    rule_number_format: Optional[str] = None
    # Функция преобразования значения после перевода в ячейку (force_int для «число, 0 знаков»)
    convert: Optional[Callable[[Any], Any]] = None


@dataclass
class SheetPlan:
    title: str
    df: pd.DataFrame
    header_styles: List[CellStyle]
    columns: List[ColumnPlan]
    widths: Dict[str, float]
    freeze: Optional[str]
    auto_filter: Optional[str]
    # Лист без колонок: прежний путь создавал пустую A1 с оформлением заголовка
    empty_header_style: Optional[CellStyle] = None


class _StyleTemplates:
    """Готовые StyleArray по ключу — присваивание стиля один раз, дальше копия массива индексов."""

    def __init__(self, ws: Any) -> None:
        self.ws = ws
        self.cache: Dict[Any, Any] = {}

    def get(self, key: Any, style: CellStyle, number_format: Optional[str]) -> Any:
        arr = self.cache.get(key)
        if arr is None:
            cell = WriteOnlyCell(self.ws)
            if style.font is not None:
                cell.font = style.font
            if style.fill is not None:
                cell.fill = style.fill
            if style.border is not None:
                cell.border = style.border
            if style.alignment is not None:
                cell.alignment = style.alignment
            if number_format is not None:
                cell.number_format = number_format
            arr = cell._style
            self.cache[key] = arr
        return arr


def _install_fast_rows(ws: Any) -> bool:
    """
    Быстрый путь записи строк: одна переиспользуемая ячейка на лист вместо объекта WriteOnlyCell на
    каждое значение (openpyxl записывает ячейку сразу, как только генератор строки её отдал).
    Строка — список пар (значение, StyleArray или None). Порядок как в Cell: значение привязывается
    к черновому стилю (openpyxl при записи даты может сам поставить формат), затем ставится шаблон
    стиля (только чтение, без копирования). Нет внутреннего метода — False, используется обычный путь.
    """
    if not hasattr(ws, "_values_to_row"):
        return False
    from openpyxl.styles.cell_style import StyleArray

    cell = WriteOnlyCell(ws)
    empty = StyleArray()

    def _values_to_row(values, row_idx):
        for col_idx, (value, style_arr) in enumerate(values, 1):
            if value is None and style_arr is None:
                continue
            cell._style = StyleArray()
            cell.value = value
            cell._style = style_arr if style_arr is not None else empty
            cell.column = col_idx
            cell.row = row_idx
            yield cell

    ws._values_to_row = _values_to_row
    return True


def write_workbook(output_path: str, plans: List[SheetPlan], active_title: Optional[str] = None) -> None:
    """Записать книгу по планам листов (порядок листов — порядок plans)."""
    wb = Workbook(write_only=True)
    for plan in plans:
        ws = wb.create_sheet(title=plan.title)
        for letter, width in plan.widths.items():
            ws.column_dimensions[letter].width = width
        if plan.freeze:
            ws.freeze_panes = plan.freeze
        if plan.auto_filter:
            ws.auto_filter.ref = plan.auto_filter
        templates = _StyleTemplates(ws)
        fast = _install_fast_rows(ws)
        df = plan.df
        n_cols = df.shape[1]

        def style_for(key: Any, style: CellStyle, number_format: Optional[str]) -> Any:
            if style.is_empty() and number_format is None:
                return None
            return templates.get(key, style, number_format)

        def emit(row_items: List[Tuple[Any, Any]]) -> None:
            if fast:
                ws.append(row_items)
                return
            out = []
            for value, arr in row_items:
                if arr is None:
                    out.append(value)
                elif value is None:
                    c = WriteOnlyCell(ws)
                    c._style = copy(arr)
                    out.append(c)
                else:
                    c = WriteOnlyCell(ws, value)
                    c._style = copy(arr)
                    out.append(c)
            ws.append(out)

        if n_cols == 0:
            if plan.empty_header_style is not None:
                # Пустая ячейка A1 с оформлением заголовка — как у прежнего пути
                emit([(None, templates.get(("h", 0), plan.empty_header_style, None))])
            continue
        header_items = []
        for j, name in enumerate(df.columns):
            value, fmt = excel_value_with_format(name)
            header_items.append((value, style_for(("h", j), plan.header_styles[j], fmt)))
        emit(header_items)

        col_plans = plan.columns
        styled = [not cp.style.is_empty() for cp in col_plans]
        rule_fmts = [cp.rule_number_format for cp in col_plans]
        converters = [cp.convert for cp in col_plans]
        col_templates: List[Dict[Any, Any]] = [dict() for _ in col_plans]
        conv = excel_value_with_format
        for row in df.itertuples(index=False, name=None):
            items = []
            append = items.append
            for j, raw in enumerate(row):
                value, fmt = conv(raw)
                fn = converters[j]
                if fn is not None:
                    value = fn(value)
                final_fmt = rule_fmts[j] if rule_fmts[j] is not None else fmt
                if not styled[j] and final_fmt is None:
                    append((value, None))
                    continue
                cache = col_templates[j]
                arr = cache.get(final_fmt)
                if arr is None:
                    arr = templates.get((j, final_fmt), col_plans[j].style, final_fmt)
                    cache[final_fmt] = arr
                append((value, arr))
            emit(items)
    if active_title is not None:
        titles = [p.title for p in plans]
        if active_title in titles:
            wb.active = titles.index(active_title)
    wb.save(output_path)


def column_letter(index_1_based: int) -> str:
    return get_column_letter(index_1_based)
