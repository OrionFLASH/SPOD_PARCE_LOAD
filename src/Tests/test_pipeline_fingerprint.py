# -*- coding: utf-8 -*-
"""Тесты отпечатка выходной книги (src/Tools/pipeline_fingerprint.py, TEST-01)."""

from __future__ import annotations

import sys
from pathlib import Path
from typing import List, Optional

import openpyxl
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.Tools.pipeline_fingerprint import (  # noqa: E402
    compare_fingerprints,
    diff_cells,
    fingerprint_workbook,
)


def _book(path: Path, header: List[str], rows: List[list], bold_col: Optional[int] = None,
          widths: Optional[dict] = None) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DATA"
    ws.append(header)
    for r in rows:
        ws.append(r)
    if bold_col is not None:
        for row in ws.iter_rows(min_row=2, min_col=bold_col, max_col=bold_col):
            row[0].font = Font(bold=True)
    for letter, w in (widths or {}).items():
        ws.column_dimensions[letter].width = w
    stat = wb.create_sheet("STAT_FILE")
    stat.append(["FILE_NAME", "DATA_UPDATE_DATE"])
    stat.append(["a.csv", str(path)])  # «время прогона» разное у каждой книги
    wb.save(path)
    return str(path)


def test_identical_books_match_and_volatile_ignored(tmp_path: Path) -> None:
    a = _book(tmp_path / "a.xlsx", ["K", "V"], [["x", 1], ["y", 2]])
    b = _book(tmp_path / "b.xlsx", ["K", "V"], [["x", 1], ["y", 2]])
    diffs, notes = compare_fingerprints(fingerprint_workbook(a), fingerprint_workbook(b))
    assert diffs == [] and notes == []


def test_column_order_is_note_unless_strict(tmp_path: Path) -> None:
    a = _book(tmp_path / "a.xlsx", ["K", "V"], [["x", 1]], widths={"A": 10, "B": 20})
    b = _book(tmp_path / "b.xlsx", ["V", "K"], [[1, "x"]], widths={"A": 20, "B": 10})
    fa, fb = fingerprint_workbook(a), fingerprint_workbook(b)
    diffs, notes = compare_fingerprints(fa, fb)
    assert diffs == []
    assert notes == ["[DATA] изменился только порядок колонок"]
    strict, _ = compare_fingerprints(fa, fb, strict_order=True)
    assert strict == ["[DATA] изменился только порядок колонок"]


def test_value_style_width_and_column_changes_detected(tmp_path: Path) -> None:
    base = fingerprint_workbook(_book(tmp_path / "a.xlsx", ["K", "V"], [["x", 1]], widths={"B": 10}))
    changed = fingerprint_workbook(
        _book(tmp_path / "b.xlsx", ["K", "V"], [["x", 2]], bold_col=1, widths={"B": 30})
    )
    diffs, _ = compare_fingerprints(base, changed)
    assert "[DATA] значения в колонках (1): ['V']" in diffs
    assert "[DATA] оформление в колонках (1): ['K']" in diffs
    assert "[DATA] ширина в колонках (1): ['V']" in diffs
    extra = fingerprint_workbook(_book(tmp_path / "c.xlsx", ["K", "V", "NEW"], [["x", 1, 0]], widths={"B": 10}))
    diffs, _ = compare_fingerprints(base, extra)
    assert "[DATA] колонки: пропали []; новые ['NEW']" in diffs


def test_diff_cells_aligns_columns_by_name(tmp_path: Path) -> None:
    a = _book(tmp_path / "a.xlsx", ["K", "V"], [["x", 1], ["y", 2]])
    b = _book(tmp_path / "b.xlsx", ["V", "K"], [[1, "x"], [3, "y"]])
    assert diff_cells(a, b, "DATA") == ["строка 3, колонка V: 2 → 3"]


def test_row_order_only_is_reported_briefly(tmp_path: Path) -> None:
    a = fingerprint_workbook(_book(tmp_path / "a.xlsx", ["K", "V"], [["x", 1], ["y", 2]]))
    b = fingerprint_workbook(_book(tmp_path / "b.xlsx", ["K", "V"], [["y", 2], ["x", 1]]))
    diffs, _ = compare_fingerprints(a, b)
    assert diffs == ["[DATA] изменился только порядок строк (набор строк тот же)"]
    c = fingerprint_workbook(_book(tmp_path / "c.xlsx", ["V", "K"], [[2, "y"], [1, "x"]]))
    diffs, notes = compare_fingerprints(a, c)
    assert diffs == ["[DATA] изменился только порядок строк (набор строк тот же)"]
    assert notes == ["[DATA] изменился только порядок колонок"]
