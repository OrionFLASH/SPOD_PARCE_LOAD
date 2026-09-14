#!/usr/bin/env python3
"""Разбор двух JSON ролей/бизнес-блоков в один Excel-файл (без внешних конфигов)."""

from __future__ import annotations

import argparse
import json
import logging
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

SCRIPT_DIR: Path = Path(__file__).resolve().parent
DEFAULT_BLOCK_JSON: Path = SCRIPT_DIR / "rolecode_business_block.json"
DEFAULT_VIEW_JSON: Path = SCRIPT_DIR / "rolecode_business_block_view.json"
DEFAULT_OUT_XLSX: Path = SCRIPT_DIR / "rolecode_business_block.xlsx"

# Порядок колонок для строк вида «число:ключ:значение»
ROW_VALUE_KEYS: list[str] = [
    "tbCode",
    "gosbCode",
    "roleCode",
    "employeeNumber",
    "permission",
    "permissionToView",
    "permissionToViewPref",
]

VALUE_TRIPLET_RE: re.Pattern[str] = re.compile(r"^(\d+):([^:]+):(.*)$", re.DOTALL)

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - [%(levelname)s] - %(message)s",
)
logger: logging.Logger = logging.getLogger("json_roles_to_excel")


def join_list(values: list[Any] | None) -> str:
    """Склеивает список в строку через запятую; пустой → ''."""
    if not values:
        return ""
    return ", ".join(str(v) for v in values)


def parse_triplet(raw: str) -> tuple[int, str, str] | None:
    """Разбирает строку «число:ключ:значение». При ошибке — None."""
    match: re.Match[str] | None = VALUE_TRIPLET_RE.match(raw)
    if match is None:
        logger.debug(
            "Пропуск нераспознанной value-строки [class: - | def: parse_triplet]: %r",
            raw,
        )
        return None
    row_num: int = int(match.group(1))
    key: str = match.group(2)
    value: str = match.group(3)
    return row_num, key, value


def parse_block_json(data: dict[str, Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """
    Из parameters/bundleList собирает:
    - meta: сведения о bundle;
    - rows: строки, сгруппированные по числу из value.
    """
    meta_rows: list[dict[str, Any]] = []
    all_rows: list[dict[str, Any]] = []

    parameters: list[dict[str, Any]] = data.get("parameters") or []
    for param in parameters:
        parameter_id: Any = param.get("parameterId")
        for bundle in param.get("bundleList") or []:
            path_parts: list[str] = []
            for item in bundle.get("path") or []:
                code: str = str(item.get("code", ""))
                value: str = str(item.get("value", ""))
                path_parts.append(f"{code}={value}")

            meta_rows.append(
                {
                    "parameterId": parameter_id,
                    "bundleId": bundle.get("id"),
                    "active": bundle.get("active"),
                    "createDate": bundle.get("createDate"),
                    "path": "; ".join(path_parts),
                    "valuesCount": len(bundle.get("values") or []),
                }
            )

            by_row: dict[int, dict[str, str]] = defaultdict(dict)
            for raw in bundle.get("values") or []:
                parsed: tuple[int, str, str] | None = parse_triplet(str(raw))
                if parsed is None:
                    continue
                row_num, key, value = parsed
                by_row[row_num][key] = value

            for row_num in sorted(by_row.keys()):
                fields: dict[str, str] = by_row[row_num]
                row: dict[str, Any] = {
                    "rowNum": row_num,
                    "parameterId": parameter_id,
                    "bundleId": bundle.get("id"),
                }
                for key in ROW_VALUE_KEYS:
                    row[key] = fields.get(key, "")
                # На случай неизвестных ключей — дописываем в конец
                for key, value in sorted(fields.items()):
                    if key not in ROW_VALUE_KEYS:
                        row[key] = value
                all_rows.append(row)

    logger.info(
        "Block JSON: meta=%s, строк=%s [def: parse_block_json]",
        len(meta_rows),
        len(all_rows),
    )
    return meta_rows, all_rows


def parse_view_json(data: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """
    Разбирает roleCodeBusinessBlockView:
    - rules: одна строка на правило (массивы склеены);
    - pairs: нормализованные пары роль ↔ бизнес-блок (+ статус).
    """
    rules: list[dict[str, Any]] = []
    pairs: list[dict[str, Any]] = []

    for idx, item in enumerate(data, start=1):
        role_codes: list[Any] = item.get("roleCode") or []
        blocks: list[Any] = item.get("businessBlockToView") or []
        statuses: list[Any] = item.get("employeeStatus") or []
        employees: list[Any] = item.get("employeeNumber") or []
        tb_codes: list[Any] = item.get("tbCode") or []
        gosb_codes: list[Any] = item.get("gosbCode") or []

        rules.append(
            {
                "ruleId": idx,
                "tbCode": join_list(tb_codes),
                "gosbCode": join_list(gosb_codes),
                "roleCode": join_list(role_codes),
                "employeeNumber": join_list(employees),
                "employeeStatus": join_list(statuses),
                "businessBlockToView": join_list(blocks),
                "rolesCount": len(role_codes),
                "blocksCount": len(blocks),
            }
        )

        # Пары: роль×блок; если ролей нет — по статусу сотрудника
        anchors: list[tuple[str, str]] = [("roleCode", str(r)) for r in role_codes]
        if not anchors and statuses:
            anchors = [("employeeStatus", str(s)) for s in statuses]
        if not anchors:
            anchors = [("roleCode", "")]

        for anchor_kind, anchor_value in anchors:
            for block in blocks or [""]:
                pairs.append(
                    {
                        "ruleId": idx,
                        "anchorKind": anchor_kind,
                        "anchorValue": anchor_value,
                        "businessBlock": str(block) if block != "" else "",
                        "tbCode": join_list(tb_codes),
                        "gosbCode": join_list(gosb_codes),
                        "employeeNumber": join_list(employees),
                    }
                )

    logger.info(
        "View JSON: правил=%s, пар=%s [def: parse_view_json]",
        len(rules),
        len(pairs),
    )
    return rules, pairs


def autosize_columns(ws: Any, max_width: int = 60) -> None:
    """Подгоняет ширину колонок по содержимому (с ограничением)."""
    for col_idx, column_cells in enumerate(ws.columns, start=1):
        length: int = 0
        for cell in column_cells:
            value: str = "" if cell.value is None else str(cell.value)
            length = max(length, len(value))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_width, max(10, length + 2))


def write_sheet(wb: Workbook, title: str, rows: list[dict[str, Any]], columns: list[str]) -> None:
    """Пишет лист с заголовком и строками словарей."""
    ws = wb.create_sheet(title)
    header_font: Font = Font(bold=True)
    for col_idx, name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, name in enumerate(columns, start=1):
            value: Any = row.get(name, "")
            if isinstance(value, bool):
                value = "true" if value else "false"
            ws.cell(row=row_idx, column=col_idx, value=value)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    autosize_columns(ws)
    logger.info("Лист «%s»: строк данных=%s [def: write_sheet]", title, len(rows))


def build_workbook(
    meta_rows: list[dict[str, Any]],
    block_rows: list[dict[str, Any]],
    view_rules: list[dict[str, Any]],
    view_pairs: list[dict[str, Any]],
) -> Workbook:
    """Собирает логичные листы Excel."""
    wb: Workbook = Workbook()
    # Удаляем дефолтный лист
    default_name: str = wb.sheetnames[0]
    wb.remove(wb[default_name])

    write_sheet(
        wb,
        "BlockMeta",
        meta_rows,
        ["parameterId", "bundleId", "active", "createDate", "path", "valuesCount"],
    )

    block_columns: list[str] = ["rowNum", "parameterId", "bundleId", *ROW_VALUE_KEYS]
    # Доп. ключи, если встретились
    extra_keys: list[str] = []
    for row in block_rows:
        for key in row:
            if key not in block_columns and key not in extra_keys:
                extra_keys.append(key)
    write_sheet(wb, "BlockRows", block_rows, block_columns + sorted(extra_keys))

    write_sheet(
        wb,
        "ViewRules",
        view_rules,
        [
            "ruleId",
            "tbCode",
            "gosbCode",
            "roleCode",
            "employeeNumber",
            "employeeStatus",
            "businessBlockToView",
            "rolesCount",
            "blocksCount",
        ],
    )

    write_sheet(
        wb,
        "ViewRoleBlock",
        view_pairs,
        [
            "ruleId",
            "anchorKind",
            "anchorValue",
            "businessBlock",
            "tbCode",
            "gosbCode",
            "employeeNumber",
        ],
    )

    return wb


def load_json(path: Path) -> Any:
    """Читает JSON из файла."""
    logger.info("Чтение JSON: %s [def: load_json]", path)
    with path.open(encoding="utf-8") as fh:
        return json.load(fh)


def parse_args() -> argparse.Namespace:
    """Аргументы CLI; пути по умолчанию — рядом со скриптом."""
    parser: argparse.ArgumentParser = argparse.ArgumentParser(
        description=(
            "Разбор rolecode_business_block.json (число:ключ:значение) "
            "и rolecode_business_block_view.json в Excel."
        )
    )
    parser.add_argument(
        "--block",
        type=Path,
        default=DEFAULT_BLOCK_JSON,
        help=f"JSON с values «число:ключ:значение» (по умолчанию: {DEFAULT_BLOCK_JSON.name})",
    )
    parser.add_argument(
        "--view",
        type=Path,
        default=DEFAULT_VIEW_JSON,
        help=f"JSON roleCodeBusinessBlockView (по умолчанию: {DEFAULT_VIEW_JSON.name})",
    )
    parser.add_argument(
        "--out",
        type=Path,
        default=DEFAULT_OUT_XLSX,
        help=f"Выходной xlsx (по умолчанию: {DEFAULT_OUT_XLSX.name})",
    )
    return parser.parse_args()


def main() -> int:
    """Точка входа."""
    args: argparse.Namespace = parse_args()
    started: datetime = datetime.now()
    logger.info("Старт конвертации [def: main]")

    block_data: dict[str, Any] = load_json(args.block)
    view_data: list[dict[str, Any]] = load_json(args.view)

    meta_rows, block_rows = parse_block_json(block_data)
    view_rules, view_pairs = parse_view_json(view_data)

    wb: Workbook = build_workbook(meta_rows, block_rows, view_rules, view_pairs)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)

    elapsed: float = (datetime.now() - started).total_seconds()
    logger.info(
        "Готово: %s (BlockRows=%s, ViewRules=%s, ViewRoleBlock=%s) за %.2f с [def: main]",
        args.out,
        len(block_rows),
        len(view_rules),
        len(view_pairs),
        elapsed,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
